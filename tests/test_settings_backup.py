from __future__ import annotations

import os
import subprocess
import uuid
from dataclasses import replace
from datetime import datetime, timedelta
from pathlib import Path

import pytest
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config.settings import get_settings
from app.email.configuration import load_smtp_config
from app.models.enums import SettingCategory, UserRole
from app.models.settings import AppSetting
from app.printing.shop_profile import load_shop_profile
from app.security.authentication import AuthenticatedUser
from app.services.backup_service import BackupService
from app.services.settings_service import SettingsService
from app.utils.exceptions import InfrastructureError, PermissionDeniedError, ValidationError


def test_encrypted_settings_validation_and_smtp_loading(
    db_session: Session, owner: AuthenticatedUser
) -> None:
    settings = get_settings()
    service = SettingsService(db_session, settings.app_secret_key.get_secret_value())
    service.set(
        actor=owner,
        category=SettingCategory.GENERAL,
        key="currency",
        value="usd",
    )
    service.set(
        actor=owner,
        category=SettingCategory.EMAIL,
        key="smtp_host",
        value="smtp.test.invalid",
    )
    service.set(
        actor=owner,
        category=SettingCategory.EMAIL,
        key="smtp_from_email",
        value="SHOP@TEST.INVALID",
    )
    secret = service.set(
        actor=owner,
        category=SettingCategory.EMAIL,
        key="smtp_password",
        value="smtp-test-secret",
        is_secret=True,
    )
    db_session.flush()
    assert service.get(SettingCategory.GENERAL, "currency") == "USD"
    assert service.get(SettingCategory.EMAIL, "smtp_password") == "smtp-test-secret"
    stored = db_session.scalar(select(AppSetting).where(AppSetting.id == secret.id))
    assert stored is not None
    assert stored.value != "smtp-test-secret"
    smtp = load_smtp_config(db_session, settings)
    assert smtp.host == "smtp.test.invalid"
    assert smtp.from_email == "shop@test.invalid"
    assert smtp.password == "smtp-test-secret"

    invalid_values = (
        (SettingCategory.GENERAL, "currency", "RUPEES"),
        (SettingCategory.GENERAL, "timezone", "Mars/Olympus"),
        (SettingCategory.GENERAL, "invoice_prefix", "bad prefix!"),
        (SettingCategory.EMAIL, "smtp_port", "70000"),
        (SettingCategory.PRINTER, "receipt_width", "72"),
        (SettingCategory.SECURITY, "session_timeout", "2"),
        (SettingCategory.BACKUP, "directory", ""),
    )
    for category, key, value in invalid_values:
        with pytest.raises(ValidationError):
            service.set(actor=owner, category=category, key=key, value=value)


def test_settings_permissions(db_session: Session, owner: AuthenticatedUser) -> None:
    salesperson = replace(owner, role=UserRole.SALESPERSON)
    with pytest.raises(PermissionDeniedError):
        SettingsService(db_session, get_settings().app_secret_key.get_secret_value()).set(
            actor=salesperson,
            category=SettingCategory.SHOP,
            key="name",
            value="No access",
        )


def test_shop_profile_fields_are_optional_and_load_for_receipts(
    db_session: Session,
    owner: AuthenticatedUser,
) -> None:
    settings = get_settings()
    service = SettingsService(db_session, settings.app_secret_key.get_secret_value())
    for key, value in {
        "name": "",
        "owner_name": "Shop Owner",
        "address": "",
        "phone": "",
        "email": "",
        "website": "example.test",
        "tax_information": "",
        "logo_path": "",
    }.items():
        service.set(
            actor=owner,
            category=SettingCategory.SHOP,
            key=key,
            value=value,
        )
    db_session.flush()

    profile = load_shop_profile(db_session, settings)

    assert profile.name == ""
    assert profile.owner_name == "Shop Owner"
    assert profile.website == "example.test"
    assert profile.address == ""
    assert profile.logo_path is None


def test_backup_success_then_retention(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    old_backup = tmp_path / "backup_2020-01-01_000000.dump"
    old_backup.write_bytes(b"old-valid-backup")
    expired = (datetime.now() - timedelta(days=90)).timestamp()
    os.utime(old_backup, (expired, expired))
    settings = get_settings().model_copy(
        update={
            "backup_directory": tmp_path,
            "backup_retention_days": 30,
            "backup_compress": True,
        }
    )
    monkeypatch.setattr("app.services.backup_service.shutil.which", lambda name: f"/{name}")

    def fake_run(command: list[str], **_kwargs: object) -> subprocess.CompletedProcess[str]:
        destination = Path(command[command.index("--file") + 1])
        destination.write_bytes(b"valid-postgresql-backup")
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr("app.services.backup_service.subprocess.run", fake_run)
    result = BackupService(settings).create_backup(uuid.uuid4())
    assert result.path.is_file()
    assert result.size_bytes > 0
    assert result.removed_expired == (old_backup,)
    assert not old_backup.exists()


def test_failed_backup_preserves_existing_files(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    old_backup = tmp_path / "backup_2020-01-01_000000.sql"
    old_backup.write_bytes(b"keep-me")
    settings = get_settings().model_copy(update={"backup_directory": tmp_path})
    monkeypatch.setattr("app.services.backup_service.shutil.which", lambda _name: "/pg_dump")

    def fail(command: list[str], **_kwargs: object) -> subprocess.CompletedProcess[str]:
        raise subprocess.CalledProcessError(1, command, stderr="simulated pg_dump failure")

    monkeypatch.setattr("app.services.backup_service.subprocess.run", fail)
    with pytest.raises(InfrastructureError, match="simulated"):
        BackupService(settings).create_backup()
    assert old_backup.read_bytes() == b"keep-me"
    assert not list(tmp_path.glob("*.part"))


def test_restore_is_owner_only_and_restricted_to_backup_directory(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    owner: AuthenticatedUser,
) -> None:
    settings = get_settings().model_copy(update={"backup_directory": tmp_path})
    service = BackupService(settings)
    backup = tmp_path / "backup_2026-08-14_210000.dump"
    backup.write_bytes(b"backup")
    salesperson = replace(owner, role=UserRole.SALESPERSON)
    with pytest.raises(PermissionDeniedError):
        service.restore(backup, salesperson)
    with pytest.raises(ValidationError, match="configured"):
        service.restore(tmp_path.parent / "outside.dump", owner)
    commands: list[list[str]] = []
    monkeypatch.setattr("app.services.backup_service.shutil.which", lambda name: f"/{name}")

    def fake_run(command: list[str], **_kwargs: object) -> subprocess.CompletedProcess[str]:
        commands.append(command)
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr("app.services.backup_service.subprocess.run", fake_run)
    service.restore(backup, owner)
    assert commands[0][0] == "/pg_restore"
    assert "--single-transaction" in commands[0]


def test_data_export_writes_every_table_and_names_the_file_for_the_period(
    db_session: Session,
    owner: AuthenticatedUser,
    customer: object,
    product: object,
    tmp_path: Path,
) -> None:
    import csv
    import io
    import zipfile

    from openpyxl import load_workbook

    from app.services.data_export_service import REDACTED_TEXT, DataExportService

    settings = get_settings()
    SettingsService(db_session, settings.app_secret_key.get_secret_value()).set(
        actor=owner,
        category=SettingCategory.EMAIL,
        key="smtp_password",
        value="smtp-test-secret",
        is_secret=True,
    )
    db_session.flush()

    result = DataExportService(db_session, settings).export(tmp_path, owner)

    assert result.workbook.is_file()
    assert result.csv_archive.is_file()
    assert result.period_start is not None
    assert result.period_end is not None
    expected = f"shop-data_from_{result.period_start:%Y-%m-%d}_to_{result.period_end:%Y-%m-%d}"
    assert result.workbook.name == f"{expected}.xlsx"
    assert result.csv_archive.name == f"{expected}.csv.zip"

    exported = {dataset.table: dataset.rows for dataset in result.datasets}
    assert exported["products"] >= 1
    assert exported["customers"] >= 1
    assert "alembic_version" not in exported

    workbook = load_workbook(result.workbook)
    assert "Products" in workbook.sheetnames
    headers = [cell.value for cell in workbook["Products"][1]]
    assert headers[0] == "Id", "the key column leads every sheet"
    assert {"Name", "Manufacturer", "Default sale price"} <= set(headers)

    with zipfile.ZipFile(result.csv_archive) as archive:
        names = set(archive.namelist())
        assert {"products.csv", "customers.csv", "sales.csv"} <= names
        users = list(
            csv.DictReader(io.StringIO(archive.read("users.csv").decode("utf-8-sig"), newline=""))
        )
        settings_rows = list(
            csv.DictReader(
                io.StringIO(archive.read("app_settings.csv").decode("utf-8-sig"), newline="")
            )
        )
    assert users, "the owner account is exported"
    assert {row["Password hash"] for row in users} == {REDACTED_TEXT}
    secrets = [row for row in settings_rows if row["Key"] == "smtp_password"]
    assert secrets and secrets[0]["Value"] == REDACTED_TEXT


def test_data_export_is_refused_to_roles_without_settings_access(
    db_session: Session, owner_model: object, tmp_path: Path
) -> None:
    from app.services.data_export_service import DataExportService

    salesperson = replace(AuthenticatedUser.from_model(owner_model), role=UserRole.SALESPERSON)  # type: ignore[arg-type]
    with pytest.raises(PermissionDeniedError):
        DataExportService(db_session, get_settings()).export(tmp_path, salesperson)
    assert list(tmp_path.iterdir()) == []


def test_a_second_data_export_of_the_same_span_keeps_the_first(
    db_session: Session, owner: AuthenticatedUser, product: object, tmp_path: Path
) -> None:
    from app.services.data_export_service import DataExportService

    service = DataExportService(db_session, get_settings())
    first = service.export(tmp_path, owner)
    second = service.export(tmp_path, owner)

    assert first.workbook.is_file() and first.csv_archive.is_file()
    assert second.workbook != first.workbook
    assert second.workbook.name.endswith("_2.xlsx")
    # Both halves of one export share a name, so a folder stays readable.
    assert second.csv_archive.name == second.workbook.name.replace(".xlsx", ".csv.zip")


def test_the_return_prefix_is_configurable_and_numbers_the_credit_note(
    db_session: Session, owner: AuthenticatedUser
) -> None:
    """The prefix was validated but had nowhere to be set; it is on General now."""

    from app.services.document_service import DocumentNumberService

    settings = get_settings()
    service = SettingsService(db_session, settings.app_secret_key.get_secret_value())
    stored = service.set(
        actor=owner, category=SettingCategory.GENERAL, key="return_prefix", value="cn"
    )
    db_session.flush()
    assert stored.value == "CN", "prefixes are normalised to upper case"
    assert service.get(SettingCategory.GENERAL, "return_prefix") == "CN"

    with pytest.raises(ValidationError):
        service.set(
            actor=owner,
            category=SettingCategory.GENERAL,
            key="return_prefix",
            value="far too long a prefix",
        )

    number = DocumentNumberService(db_session).next_number(
        "sale_return",
        service.get(SettingCategory.GENERAL, "return_prefix") or "RET",
        datetime.now(),
    )
    assert number.startswith("CN-")


@pytest.mark.ui
def test_the_settings_screen_round_trips_both_document_prefixes(
    qtbot: object, database_engine: object, owner: AuthenticatedUser
) -> None:
    from sqlalchemy.orm import sessionmaker

    from app.ui.settings.screen import SettingsScreen

    factory = sessionmaker[Session](bind=database_engine, expire_on_commit=False, autoflush=False)  # type: ignore[arg-type]
    screen = SettingsScreen(factory, owner, get_settings())
    qtbot.addWidget(screen)  # type: ignore[attr-defined]
    assert screen.invoice_prefix.text() == "INV"
    assert screen.return_prefix.text() == "RET"


def test_a_shop_name_is_refused_as_an_smtp_username(
    db_session: Session, owner: AuthenticatedUser
) -> None:
    """Gmail reports this only as a bad password, so it is caught at entry instead."""

    settings = get_settings()
    service = SettingsService(db_session, settings.app_secret_key.get_secret_value())

    with pytest.raises(ValidationError, match="cannot contain spaces"):
        service.set(
            actor=owner,
            category=SettingCategory.EMAIL,
            key="smtp_username",
            value="Khan Zari Services",
        )

    stored = service.set(
        actor=owner,
        category=SettingCategory.EMAIL,
        key="smtp_username",
        value="  shop@gmail.com  ",
    )
    db_session.flush()
    assert stored.value == "shop@gmail.com", "surrounding whitespace is trimmed, not rejected"


def test_owner_emails_are_stored_as_a_list(db_session: Session, owner: AuthenticatedUser) -> None:
    """Several people are copied on every invoice, so the one key holds a list."""

    from app.email.configuration import load_owner_emails, parse_owner_emails

    settings = get_settings()
    service = SettingsService(db_session, settings.app_secret_key.get_secret_value())
    stored = service.set(
        actor=owner,
        category=SettingCategory.EMAIL,
        key="owner_email",
        value="Owner@Shop.com, partner@shop.com; owner@shop.com",
    )
    db_session.flush()
    assert stored.value == "owner@shop.com, partner@shop.com", "duplicates are dropped"
    assert load_owner_emails(db_session, settings) == ("owner@shop.com", "partner@shop.com")

    # A shop that set one address before this was a list keeps working untouched.
    assert parse_owner_emails("solo@shop.com") == ("solo@shop.com",)

    with pytest.raises(ValidationError, match="valid email"):
        service.set(
            actor=owner, category=SettingCategory.EMAIL, key="owner_email", value="not-an-email"
        )

    # Clearing the list is allowed; nobody is copied.
    cleared = service.set(actor=owner, category=SettingCategory.EMAIL, key="owner_email", value="")
    db_session.flush()
    assert cleared.value == ""
    assert load_owner_emails(db_session, settings) == ()


@pytest.mark.ui
def test_the_settings_screen_adds_and_removes_owner_emails(
    qtbot: object, database_engine: object, owner: AuthenticatedUser
) -> None:
    from sqlalchemy.orm import sessionmaker

    from app.ui.settings.screen import SettingsScreen

    factory = sessionmaker[Session](bind=database_engine, expire_on_commit=False, autoflush=False)  # type: ignore[arg-type]
    screen = SettingsScreen(factory, owner, get_settings())
    qtbot.addWidget(screen)  # type: ignore[attr-defined]
    assert screen._owner_email_addresses() == ()
    assert screen.remove_owner_email.isEnabled() is False

    screen.owner_email_entry.setText("  Owner@Shop.com ")
    screen._add_owner_email()
    screen.owner_email_entry.setText("partner@shop.com")
    screen._add_owner_email()
    assert screen._owner_email_addresses() == ("owner@shop.com", "partner@shop.com")
    assert screen.owner_email_entry.text() == "", "the box clears ready for the next one"

    screen.owner_emails.setCurrentRow(0)
    assert screen.remove_owner_email.isEnabled() is True
    screen._remove_owner_email()
    assert screen._owner_email_addresses() == ("partner@shop.com",)
