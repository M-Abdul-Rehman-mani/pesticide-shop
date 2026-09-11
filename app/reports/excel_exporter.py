"""Professional spreadsheet export using openpyxl."""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


class ExcelExporter:
    HEADER_FILL = PatternFill("solid", fgColor="17324D")
    SUMMARY_FILL = PatternFill("solid", fgColor="DDEBF7")

    @staticmethod
    def _cell_value(
        value: object,
    ) -> bool | int | Decimal | str | datetime | date | None:
        """Return an Excel-safe scalar while retaining a real date cell.

        Excel has no timezone-aware datetime representation. Database timestamps are
        normalized to UTC before the timezone marker is removed so exports remain
        deterministic across workstations.
        """

        if isinstance(value, datetime) and value.tzinfo is not None:
            return value.astimezone(UTC).replace(tzinfo=None)
        if isinstance(value, (bool, int, Decimal, str, datetime, date)) or value is None:
            return value
        return str(value)

    def export_workbook(
        self,
        path: Path,
        datasets: Iterable[tuple[str, Sequence[str], Iterable[Sequence[object]]]],
    ) -> Path:
        """Write several datasets to one workbook, a sheet each.

        Used by the whole-database export, where the point is that one file holds
        every table rather than one report.
        """

        workbook = Workbook()
        default = workbook.active
        used: set[str] = set()
        for title, headers, rows in datasets:
            sheet = workbook.create_sheet(self._sheet_title(title, used))
            self._write_sheet(sheet, headers, rows)
        if default is not None and not workbook.sheetnames[1:]:
            # An empty workbook is still a valid file; keep the default sheet so
            # the export opens rather than erroring.
            assert isinstance(default, Worksheet)
            default.title = "Empty"
        elif default is not None:
            workbook.remove(default)
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        return path

    @staticmethod
    def _sheet_title(title: str, used: set[str]) -> str:
        # Excel limits sheet names to 31 characters and rejects duplicates.
        candidate = title[:31] or "Sheet"
        suffix = 2
        while candidate.lower() in used:
            candidate = f"{title[:28]}_{suffix}"
            suffix += 1
        used.add(candidate.lower())
        return candidate

    def _write_sheet(
        self,
        sheet: Worksheet,
        headers: Sequence[str],
        rows: Iterable[Sequence[object]],
    ) -> None:
        sheet.freeze_panes = "A2"
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(1, column, header)
            cell.fill = self.HEADER_FILL
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center")
        widths = [len(header) for header in headers]
        row_number = 1
        for row_number, values in enumerate(rows, start=2):
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row_number, column, self._cell_value(value))
                if isinstance(value, (date, datetime)):
                    cell.number_format = "dd-mmm-yyyy hh:mm"
                # Sampling keeps a large table from being measured row by row.
                if column <= len(widths) and row_number <= 200:
                    widths[column - 1] = max(widths[column - 1], len(str(cell.value or "")))
        if headers:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, row_number)}"
        for column, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = min(50, width + 2)

    def export(
        self,
        path: Path,
        *,
        title: str,
        headers: Sequence[str],
        rows: Iterable[Sequence[object]],
        summary: Sequence[object] | None = None,
        currency_columns: frozenset[int] = frozenset(),
        date_columns: frozenset[int] = frozenset(),
    ) -> Path:
        workbook = Workbook()
        sheet = workbook.active
        assert isinstance(sheet, Worksheet)
        sheet.title = title[:31]
        sheet.freeze_panes = "A3"
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        title_cell = sheet.cell(1, 1, title)
        title_cell.font = Font(size=16, bold=True, color="17324D")
        title_cell.alignment = Alignment(horizontal="center")
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(2, column, header)
            cell.fill = self.HEADER_FILL
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center")
        row_number = 3
        for values in rows:
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row_number, column, self._cell_value(value))
                if column in currency_columns and isinstance(value, (Decimal, int)):
                    cell.number_format = "#,##0.00;[Red]-#,##0.00"
                elif column in date_columns and isinstance(value, (date, datetime)):
                    cell.number_format = "dd-mmm-yyyy hh:mm"
            row_number += 1
        if summary:
            for column, value in enumerate(summary, start=1):
                cell = sheet.cell(row_number, column, self._cell_value(value))
                cell.fill = self.SUMMARY_FILL
                cell.font = Font(bold=True)
                if column in currency_columns:
                    cell.number_format = "#,##0.00;[Red]-#,##0.00"
        last_row = row_number if summary else max(2, row_number - 1)
        sheet.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{last_row}"
        for column, header in enumerate(headers, start=1):
            display_values = [
                str(sheet.cell(row, column).value or "") for row in range(2, last_row + 1)
            ]
            width = min(50, max(len(header), *(len(value) for value in display_values)) + 2)
            sheet.column_dimensions[get_column_letter(column)].width = width
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        return path
